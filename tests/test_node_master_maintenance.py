from openpyxl import Workbook
import pytest

from src.data.node_master_maintenance import (
    NODE_MASTER_HEADERS,
    NodeMasterMaintenanceError,
    load_node_master_maintenance_entries,
)


def test_node_master_cleans_alias_placeholders_and_ignores_other_sheet(tmp_path):
    path = tmp_path / "节点信息维护0731.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "节点信息维护"
    sheet.append(NODE_MASTER_HEADERS)
    sheet.append(
        (
            "物流节点",
            "港口码头",
            "海港、内河码头",
            "散粮、集装箱",
            "2000",
            "测试经营部",
            "测试港",
            "24",
            "/",
            "测试港别名",
            "广东省",
            "广州市",
            "黄埔区",
            "港前路",
            "1号",
            113.5,
            23.0,
        )
    )
    other = workbook.create_sheet("Sheet1")
    other.append(NODE_MASTER_HEADERS)
    other.append(
        (
            "物流节点",
            "港口码头",
            "内河码头",
            "散粮",
            "/",
            "",
            "不应加载",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            110,
            22,
        )
    )
    workbook.save(path)

    entries = load_node_master_maintenance_entries(path)

    assert len(entries) == 1
    assert entries[0].aliases == ("测试港别名",)
    assert entries[0].port_attributes == ("海港", "内河码头")
    assert entries[0].package_types == ("散粮", "集装箱")
    assert entries[0].draft_capacity == "2000"
    assert entries[0].has_seaport_attribute
    assert entries[0].has_inland_port_attribute


def test_node_master_merges_customer_multi_role_into_one_node(tmp_path):
    path = tmp_path / "节点信息维护0731.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "节点信息维护"
    sheet.append(NODE_MASTER_HEADERS)
    sheet.append(
        (
            "客户",
            "饲料客户",
            "客户仓库",
            "散粮",
            "/",
            "经营部A",
            "测试客户有限公司",
            "",
            "",
            "",
            "广东省",
            "东莞市",
            "麻涌镇",
            "",
            "客户入口",
            113.50,
            23.00,
        )
    )
    sheet.append(
        (
            "物流节点",
            "港口码头",
            "客户仓库、内河码头",
            "散粮、集装箱",
            "2000",
            "经营部B",
            "测试客户有限公司",
            "测试客户作业点",
            "",
            "",
            "广东省",
            "东莞市",
            "麻涌镇",
            "",
            "内河作业点",
            113.51,
            23.01,
        )
    )
    workbook.save(path)

    entries = load_node_master_maintenance_entries(path)

    assert len(entries) == 1
    entry = entries[0]
    assert entry.is_customer_node
    assert entry.is_logistics_node
    assert entry.longitude == 113.51
    assert entry.latitude == 23.01
    assert entry.port_attributes == ("客户仓库", "内河码头")
    assert entry.package_types == ("散粮", "集装箱")
    assert entry.aliases == ("测试客户作业点",)
    assert "duplicate_role_logistics_coordinate_preferred" in entry.source


def test_node_master_does_not_merge_distinct_logistics_ports_with_same_name(
    tmp_path,
):
    path = tmp_path / "节点信息维护0731.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "节点信息维护"
    sheet.append(NODE_MASTER_HEADERS)
    base = [
        "物流节点",
        "港口码头",
        "内河码头",
        "散粮",
        "2000",
        "",
        "重名码头",
        "",
        "",
        "",
        "广东省",
        "广州市",
        "",
        "",
        "",
        113.0,
        23.0,
    ]
    sheet.append(base)
    other = list(base)
    other[-2] = 114.0
    other[-1] = 24.0
    sheet.append(other)
    workbook.save(path)

    with pytest.raises(NodeMasterMaintenanceError, match="坐标明显分离"):
        load_node_master_maintenance_entries(path)
