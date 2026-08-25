from __future__ import annotations

import os
from datetime import date, datetime
from pathlib import Path
from typing import Any


FREIGHT_WORKBOOK_FILE = "运费数据.xlsx"
FREIGHT_WORKBOOK_PATH_ENV = "FREIGHT_WORKBOOK_PATH"
FREIGHT_WORKBOOK_SHEET = "运价表"
FREIGHT_WORKBOOK_HEADERS = (
    "始发",
    "到达",
    "运输方式",
    "包装方式",
    "适用品种",
    "费用",
    "费用单位",
    "价格来源",
    "维护日期",
)


class FreightWorkbookError(ValueError):
    """Raised when the maintained freight workbook is unsafe to consume."""


def find_freight_workbook(data_dir: Path) -> Path | None:
    configured_path = os.environ.get(FREIGHT_WORKBOOK_PATH_ENV, "").strip()
    if configured_path:
        path = Path(configured_path)
        if not path.is_absolute():
            path = data_dir / path
        if not path.exists():
            raise FreightWorkbookError(
                f"{FREIGHT_WORKBOOK_PATH_ENV} 指向的运费工作簿不存在：{path}"
            )
        if not path.is_file():
            raise FreightWorkbookError(
                f"{FREIGHT_WORKBOOK_PATH_ENV} 必须指向文件：{path}"
            )
        if path.suffix.lower() != ".xlsx":
            raise FreightWorkbookError(
                f"{FREIGHT_WORKBOOK_PATH_ENV} 必须指向 .xlsx 工作簿：{path}"
            )
        return path
    matches = sorted(data_dir.rglob(FREIGHT_WORKBOOK_FILE))
    if len(matches) > 1:
        joined = "; ".join(str(path) for path in matches)
        raise FreightWorkbookError(
            f"DATA_DIR 下存在多个 {FREIGHT_WORKBOOK_FILE}：{joined}"
        )
    return matches[0] if matches else None


def load_freight_workbook_rows(
    path: Path,
) -> tuple[tuple[int, dict[str, Any]], ...]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FreightWorkbookError(
            "读取运费数据工作簿需要本地 openpyxl 依赖。"
        ) from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise FreightWorkbookError(
            f"{path.name} 无法作为 Excel 工作簿读取：{exc}"
        ) from exc
    try:
        if FREIGHT_WORKBOOK_SHEET not in workbook.sheetnames:
            raise FreightWorkbookError(
                f"{path.name} 缺少工作表：{FREIGHT_WORKBOOK_SHEET}"
            )
        sheet = workbook[FREIGHT_WORKBOOK_SHEET]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise FreightWorkbookError(
                f"{path.name}/{FREIGHT_WORKBOOK_SHEET} 为空。"
            ) from exc
        headers = tuple(
            str(value).strip() if value is not None else ""
            for value in raw_headers
        )
        if headers != FREIGHT_WORKBOOK_HEADERS:
            raise FreightWorkbookError(
                f"{path.name}/{FREIGHT_WORKBOOK_SHEET} 表头必须严格为："
                + "、".join(FREIGHT_WORKBOOK_HEADERS)
            )

        result: list[tuple[int, dict[str, Any]]] = []
        for excel_row, values in enumerate(rows, start=2):
            if all(_is_blank(value) for value in values):
                continue
            row = dict(zip(headers, values, strict=True))
            row["维护日期"] = _normalized_maintenance_date(
                row.get("维护日期"),
                path=path,
                excel_row=excel_row,
            )
            result.append((excel_row, row))
        if not result:
            raise FreightWorkbookError(
                f"{path.name}/{FREIGHT_WORKBOOK_SHEET} 没有运价记录。"
            )
        return tuple(result)
    finally:
        workbook.close()


def _normalized_maintenance_date(
    value: object,
    *,
    path: Path,
    excel_row: int,
) -> str | None:
    if _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    raise FreightWorkbookError(
        f"{path.name}/{FREIGHT_WORKBOOK_SHEET} 第 {excel_row} 行维护日期"
        "必须为日期、YYYY-MM-DD 文本或空值。"
    )


def _is_blank(value: object) -> bool:
    return value is None or (
        isinstance(value, str) and not value.strip()
    )
