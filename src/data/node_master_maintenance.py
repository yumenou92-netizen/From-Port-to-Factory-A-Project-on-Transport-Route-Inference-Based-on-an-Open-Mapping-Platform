from __future__ import annotations

import math
import re
from dataclasses import dataclass, replace
from pathlib import Path


NODE_MASTER_MAINTENANCE_FILE = "节点信息维护0731.xlsx"
NODE_MASTER_SHEET_NAME = "节点信息维护"
NODE_MASTER_HEADERS = (
    "节点类型",
    "节点性质",
    "码头属性",
    "包装方式",
    "吃水能力（仅码头/码头）",
    "经营部",
    "全称",
    "简称1",
    "简称2",
    "简称3",
    "省/自治区/直辖市",
    "地级市",
    "区/县/镇",
    "街道",
    "详细地址",
    "经度",
    "纬度",
)
EMPTY_ALIAS_MARKERS = frozenset({"24", "/"})
DUPLICATE_COORDINATE_TOLERANCE = 0.02


class NodeMasterMaintenanceError(ValueError):
    """Raised when the maintained node-master workbook is unsafe to consume."""


@dataclass(frozen=True)
class NodeMasterMaintenanceEntry:
    row_number: int
    node_type: str
    node_nature: str
    port_attributes: tuple[str, ...]
    package_types: tuple[str, ...]
    draft_capacity: str | None
    business_unit: str | None
    full_name: str
    aliases: tuple[str, ...]
    province: str | None
    city: str | None
    district: str | None
    address: str | None
    longitude: float
    latitude: float
    source: str

    @property
    def all_names(self) -> tuple[str, ...]:
        return (self.full_name, *self.aliases)

    @property
    def is_logistics_node(self) -> bool:
        return "物流节点" in self.node_type

    @property
    def is_customer_node(self) -> bool:
        return "客户" in self.node_type

    @property
    def is_port_facility(self) -> bool:
        return "港口码头" in self.node_nature or any(
            value in {"海港", "内河码头"} for value in self.port_attributes
        )

    @property
    def has_seaport_attribute(self) -> bool:
        return "海港" in self.port_attributes

    @property
    def has_inland_port_attribute(self) -> bool:
        return "内河码头" in self.port_attributes

    @property
    def has_rail_attribute(self) -> bool:
        return (
            "铁路站点" in self.node_nature
            or "铁路专用线" in self.port_attributes
        )

    @property
    def location_text(self) -> str:
        return "；".join(
            value
            for value in (
                self.full_name,
                *self.aliases,
                self.province,
                self.city,
                self.district,
            )
            if value
        )


def find_node_master_maintenance_file(data_dir: Path) -> Path | None:
    matches = sorted(data_dir.rglob(NODE_MASTER_MAINTENANCE_FILE))
    if len(matches) > 1:
        joined = "; ".join(str(path) for path in matches)
        raise NodeMasterMaintenanceError(
            f"DATA_DIR 下存在多个 {NODE_MASTER_MAINTENANCE_FILE}：{joined}"
        )
    return matches[0] if matches else None


def load_node_master_maintenance_entries(
    path: Path,
) -> tuple[NodeMasterMaintenanceEntry, ...]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise NodeMasterMaintenanceError(
            "读取节点信息维护表需要本地 openpyxl 依赖。"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if NODE_MASTER_SHEET_NAME not in workbook.sheetnames:
            raise NodeMasterMaintenanceError(
                f"{path.name} 缺少工作表 {NODE_MASTER_SHEET_NAME}。"
            )
        worksheet = workbook[NODE_MASTER_SHEET_NAME]
        rows = worksheet.iter_rows(values_only=True)
        try:
            first_row = next(rows)
        except StopIteration as exc:
            raise NodeMasterMaintenanceError(f"{path.name} 是空工作簿。") from exc
        header = tuple(_clean_text(value) for value in first_row)
        if header[: len(NODE_MASTER_HEADERS)] != NODE_MASTER_HEADERS:
            expected = "、".join(NODE_MASTER_HEADERS)
            actual = "、".join(
                value or "<空>"
                for value in header[: len(NODE_MASTER_HEADERS)]
            )
            raise NodeMasterMaintenanceError(
                f"{path.name} 表头必须为 {expected}；实际为 {actual}。"
            )

        entries: list[NodeMasterMaintenanceEntry] = []
        seen_by_name: dict[str, int] = {}
        for row_number, row in enumerate(rows, start=2):
            values = tuple(row[: len(NODE_MASTER_HEADERS)]) + (None,) * max(
                0,
                len(NODE_MASTER_HEADERS) - len(row),
            )
            if not any(value is not None and _clean_text(value) for value in values):
                continue
            (
                node_type,
                node_nature,
                port_attributes,
                package_types,
                draft_capacity,
                business_unit,
                full_name,
                short_name_1,
                short_name_2,
                short_name_3,
                province,
                city,
                district,
                street,
                detailed_address,
                longitude,
                latitude,
            ) = values[: len(NODE_MASTER_HEADERS)]

            clean_full_name = _required_text(
                full_name, path, row_number, "全称"
            )
            aliases = tuple(
                dict.fromkeys(
                    alias
                    for alias in (
                        _alias_text(short_name_1),
                        _alias_text(short_name_2),
                        _alias_text(short_name_3),
                    )
                    if alias and alias != clean_full_name
                )
            )
            address = "".join(
                value
                for value in (
                    _optional_text(street),
                    _optional_text(detailed_address),
                )
                if value
            )
            entry = NodeMasterMaintenanceEntry(
                row_number=row_number,
                node_type=_required_text(
                    node_type, path, row_number, "节点类型"
                ),
                node_nature=_required_text(
                    node_nature, path, row_number, "节点性质"
                ),
                port_attributes=_split_text(port_attributes),
                package_types=_split_text(package_types),
                draft_capacity=_marker_aware_optional_text(draft_capacity),
                business_unit=_optional_text(business_unit),
                full_name=clean_full_name,
                aliases=aliases,
                province=_optional_text(province),
                city=_optional_text(city),
                district=_optional_text(district),
                address=address or None,
                longitude=_coordinate(
                    longitude, path, row_number, "经度", -180, 180
                ),
                latitude=_coordinate(
                    latitude, path, row_number, "纬度", -90, 90
                ),
                source=f"{path.name}#{NODE_MASTER_SHEET_NAME}!{row_number}",
            )
            previous_index = seen_by_name.get(clean_full_name)
            if previous_index is not None:
                entries[previous_index] = _merge_duplicate_records(
                    entries[previous_index],
                    entry,
                    path=path,
                )
                continue
            seen_by_name[clean_full_name] = len(entries)
            entries.append(entry)
        return tuple(entries)
    finally:
        workbook.close()


def _same_business_record(
    left: NodeMasterMaintenanceEntry,
    right: NodeMasterMaintenanceEntry,
) -> bool:
    return (
        left.node_type,
        left.node_nature,
        left.port_attributes,
        left.package_types,
        left.draft_capacity,
        left.business_unit,
        left.full_name,
        left.aliases,
        left.province,
        left.city,
        left.district,
        left.address,
        left.longitude,
        left.latitude,
    ) == (
        right.node_type,
        right.node_nature,
        right.port_attributes,
        right.package_types,
        right.draft_capacity,
        right.business_unit,
        right.full_name,
        right.aliases,
        right.province,
        right.city,
        right.district,
        right.address,
        right.longitude,
        right.latitude,
    )


def _merge_duplicate_records(
    left: NodeMasterMaintenanceEntry,
    right: NodeMasterMaintenanceEntry,
    *,
    path: Path,
) -> NodeMasterMaintenanceEntry:
    if _same_business_record(left, right):
        return replace(left, source=f"{left.source};{right.source}")
    if left.node_type != right.node_type:
        logistics_records = [
            record
            for record in (left, right)
            if record.node_type == "物流节点"
        ]
        if len(logistics_records) == 1:
            preferred = logistics_records[0]
            other = right if preferred is left else left
            return _merge_record_attributes(
                preferred,
                other,
                node_type="、".join(
                    dict.fromkeys(
                        (preferred.node_type, other.node_type)
                    )
                ),
                source_marker="duplicate_role_logistics_coordinate_preferred",
            )
        raise NodeMasterMaintenanceError(
            f"{path.name} 的全称 {left.full_name} 在第 "
            f"{left.row_number}、{right.row_number} 行分别属于 "
            f"{left.node_type}/{right.node_type}，不能自动合并。"
        )
    if (
        abs(left.longitude - right.longitude)
        > DUPLICATE_COORDINATE_TOLERANCE
        or abs(left.latitude - right.latitude)
        > DUPLICATE_COORDINATE_TOLERANCE
    ):
        if left.node_type == right.node_type == "客户":
            preferred = (
                left
                if "饲料客户" in left.node_nature
                else (
                    right
                    if "饲料客户" in right.node_nature
                    else left
                )
            )
            other = right if preferred is left else left
            return _merge_record_attributes(
                preferred,
                other,
                node_type=preferred.node_type,
                source_marker=(
                    "duplicate_customer_secondary_coordinate_ignored"
                ),
            )
        raise NodeMasterMaintenanceError(
            f"{path.name} 的全称 {left.full_name} 在第 "
            f"{left.row_number}、{right.row_number} 行坐标明显分离，"
            "不能按重复项自动合并。"
        )
    return _merge_record_attributes(
        left,
        right,
        node_type=left.node_type,
        source_marker="duplicate_record_attributes_merged",
    )


def _merge_record_attributes(
    preferred: NodeMasterMaintenanceEntry,
    other: NodeMasterMaintenanceEntry,
    *,
    node_type: str,
    source_marker: str,
) -> NodeMasterMaintenanceEntry:
    return replace(
        preferred,
        node_type=node_type,
        node_nature="、".join(
            dict.fromkeys(
                (preferred.node_nature, other.node_nature)
            )
        ),
        port_attributes=tuple(
            dict.fromkeys(
                (*preferred.port_attributes, *other.port_attributes)
            )
        ),
        package_types=tuple(
            dict.fromkeys(
                (*preferred.package_types, *other.package_types)
            )
        ),
        aliases=tuple(
            dict.fromkeys((*preferred.aliases, *other.aliases))
        ),
        draft_capacity=preferred.draft_capacity or other.draft_capacity,
        business_unit=_join_optional(
            preferred.business_unit, other.business_unit
        ),
        province=preferred.province or other.province,
        city=preferred.city or other.city,
        district=preferred.district or other.district,
        address=preferred.address or other.address,
        source=(
            f"{preferred.source};{source_marker};{other.source}"
        ),
    )


def _join_optional(left: str | None, right: str | None) -> str | None:
    values = tuple(
        dict.fromkeys(value for value in (left, right) if value)
    )
    return "、".join(values) or None


def _required_text(
    value: object,
    path: Path,
    row_number: int,
    field_name: str,
) -> str:
    text = _clean_text(value)
    if not text:
        raise NodeMasterMaintenanceError(
            f"{path.name} 第 {row_number} 行缺少必填字段：{field_name}"
        )
    return text


def _optional_text(value: object) -> str | None:
    return _clean_text(value) or None


def _marker_aware_optional_text(value: object) -> str | None:
    text = _clean_text(value)
    return None if not text or text in EMPTY_ALIAS_MARKERS else text


def _alias_text(value: object) -> str | None:
    text = _clean_text(value)
    return None if not text or text in EMPTY_ALIAS_MARKERS else text


def _clean_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _split_text(value: object) -> tuple[str, ...]:
    text = _clean_text(value)
    if not text or text in EMPTY_ALIAS_MARKERS:
        return ()
    return tuple(
        dict.fromkeys(
            item.strip()
            for item in re.split(r"[,，、;/；]+", text)
            if item.strip()
        )
    )


def _coordinate(
    value: object,
    path: Path,
    row_number: int,
    field_name: str,
    minimum: float,
    maximum: float,
) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise NodeMasterMaintenanceError(
            f"{path.name} 第 {row_number} 行 {field_name} 不是有效数值：{value}"
        ) from exc
    if not math.isfinite(number) or not minimum <= number <= maximum:
        raise NodeMasterMaintenanceError(
            f"{path.name} 第 {row_number} 行 {field_name} 超出有效范围：{value}"
        )
    return number
