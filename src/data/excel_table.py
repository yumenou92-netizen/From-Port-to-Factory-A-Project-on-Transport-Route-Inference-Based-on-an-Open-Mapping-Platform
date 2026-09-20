"""Generic, read-only local Excel table access.

The routing model never depends on a workbook's physical column spelling.
Each local source is registered through a small JSON manifest which maps its
business fields to one or more workbook headers.  This module intentionally
does not interpret freight rules; it only returns source-traceable rows.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping, Sequence


class LocalExcelTableError(ValueError):
    """Raised when a local workbook source cannot be safely read."""


@dataclass(frozen=True)
class ExcelTableSource:
    """One locally maintained workbook table and its canonical field mapping."""

    source_id: str
    table_type: str
    workbook: str
    sheet_name: str
    field_mapping: Mapping[str, tuple[str, ...]]
    defaults: Mapping[str, str] | None = None

    def __post_init__(self) -> None:
        for field_name in ("source_id", "table_type", "workbook", "sheet_name"):
            value = str(getattr(self, field_name)).strip()
            if not value:
                raise LocalExcelTableError(f"{field_name} 不能为空。")
            object.__setattr__(self, field_name, value)
        normalized_mapping: dict[str, tuple[str, ...]] = {}
        for canonical, headers in self.field_mapping.items():
            key = str(canonical).strip()
            values = tuple(str(header).strip() for header in headers if str(header).strip())
            if not key or not values:
                raise LocalExcelTableError("字段映射中的规范字段和表头候选均不能为空。")
            normalized_mapping[key] = values
        object.__setattr__(self, "field_mapping", normalized_mapping)
        object.__setattr__(
            self,
            "defaults",
            {str(key).strip(): str(value).strip() for key, value in (self.defaults or {}).items() if str(key).strip()},
        )


@dataclass(frozen=True)
class ExcelTableRow:
    """A canonicalized source row retaining source identity and Excel row."""

    source_id: str
    table_type: str
    workbook: Path
    sheet_name: str
    row_number: int
    values: Mapping[str, object]

    @property
    def source_reference(self) -> str:
        return f"{self.workbook.name}/{self.sheet_name}/第{self.row_number}行"

    def value(self, field_name: str) -> object | None:
        return self.values.get(field_name)


def load_excel_source_manifest(path: str | Path) -> tuple[ExcelTableSource, ...]:
    """Load local-only source registration without reading business tables yet."""

    manifest_path = Path(path)
    try:
        payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise LocalExcelTableError(f"未找到本地 Excel 数据源清单：{manifest_path}") from exc
    except json.JSONDecodeError as exc:
        raise LocalExcelTableError(f"Excel 数据源清单不是有效 JSON：{manifest_path.name}") from exc
    if not isinstance(payload, dict) or not isinstance(payload.get("sources"), list):
        raise LocalExcelTableError("Excel 数据源清单必须包含 sources 数组。")
    sources: list[ExcelTableSource] = []
    source_ids: set[str] = set()
    for index, item in enumerate(payload["sources"], start=1):
        if not isinstance(item, dict):
            raise LocalExcelTableError(f"sources 第 {index} 项必须是对象。")
        mapping = _parse_mapping(item.get("field_mapping"), index)
        source = ExcelTableSource(
            source_id=str(item.get("source_id", "")),
            table_type=str(item.get("table_type", "")),
            workbook=str(item.get("workbook", "")),
            sheet_name=str(item.get("sheet_name", "")),
            field_mapping=mapping,
            defaults=item.get("defaults") if isinstance(item.get("defaults"), dict) else None,
        )
        if source.source_id in source_ids:
            raise LocalExcelTableError(f"Excel 数据源清单存在重复 source_id：{source.source_id}")
        source_ids.add(source.source_id)
        sources.append(source)
    if not sources:
        raise LocalExcelTableError("Excel 数据源清单不能为空。")
    return tuple(sources)


def read_excel_table(
    source: ExcelTableSource,
    *,
    base_dir: str | Path,
    required_fields: Sequence[str],
) -> tuple[ExcelTableRow, ...]:
    """Read one sheet once, map headers, and return non-empty canonical rows."""

    workbook_path = _resolve_workbook(source.workbook, Path(base_dir))
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment boundary
        raise LocalExcelTableError("读取本地 Excel 数据源需要 openpyxl 依赖。") from exc
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception as exc:
        raise LocalExcelTableError(f"{workbook_path.name} 无法作为 Excel 工作簿读取：{exc}") from exc
    try:
        if source.sheet_name not in workbook.sheetnames:
            raise LocalExcelTableError(f"{workbook_path.name} 缺少工作表：{source.sheet_name}")
        rows = workbook[source.sheet_name].iter_rows(values_only=True)
        try:
            headers = tuple(_text(value) for value in next(rows))
        except StopIteration as exc:
            raise LocalExcelTableError(f"{workbook_path.name}/{source.sheet_name} 为空。") from exc
        positions = _resolve_positions(source, headers, required_fields)
        result: list[ExcelTableRow] = []
        for row_number, values in enumerate(rows, start=2):
            mapped = dict(source.defaults or {})
            for field_name, position in positions.items():
                value = values[position] if position < len(values) else None
                if value is not None and (not isinstance(value, str) or value.strip()):
                    mapped[field_name] = value.strip() if isinstance(value, str) else value
            if any(value not in (None, "") for value in mapped.values()):
                result.append(
                    ExcelTableRow(
                        source.source_id,
                        source.table_type,
                        workbook_path,
                        source.sheet_name,
                        row_number,
                        mapped,
                    )
                )
        return tuple(result)
    finally:
        workbook.close()


def _parse_mapping(value: object, index: int) -> dict[str, tuple[str, ...]]:
    if not isinstance(value, dict):
        raise LocalExcelTableError(f"sources 第 {index} 项缺少 field_mapping 对象。")
    mapping: dict[str, tuple[str, ...]] = {}
    for canonical, headers in value.items():
        if isinstance(headers, str):
            mapping[str(canonical)] = (headers,)
        elif isinstance(headers, list):
            mapping[str(canonical)] = tuple(str(header) for header in headers)
        else:
            raise LocalExcelTableError(f"sources 第 {index} 项字段 {canonical} 的表头映射必须是字符串或数组。")
    return mapping


def _resolve_workbook(workbook: str, base_dir: Path) -> Path:
    candidate = Path(workbook)
    resolved = candidate if candidate.is_absolute() else base_dir / candidate
    if not resolved.is_file():
        raise LocalExcelTableError(f"未找到数据源工作簿：{resolved}")
    return resolved


def _resolve_positions(
    source: ExcelTableSource,
    headers: tuple[str, ...],
    required_fields: Sequence[str],
) -> dict[str, int]:
    lookup = {header: index for index, header in enumerate(headers) if header}
    positions: dict[str, int] = {}
    missing: list[str] = []
    for field_name, aliases in source.field_mapping.items():
        position = next((lookup[alias] for alias in aliases if alias in lookup), None)
        if position is not None:
            positions[field_name] = position
    for field_name in required_fields:
        if field_name not in positions and not (source.defaults or {}).get(field_name):
            aliases = source.field_mapping.get(field_name, ())
            missing.append(f"{field_name} ({' / '.join(aliases)})")
    if missing:
        raise LocalExcelTableError(
            f"{source.source_id} 缺少必需字段映射或表头：" + "；".join(missing)
        )
    return positions


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""
