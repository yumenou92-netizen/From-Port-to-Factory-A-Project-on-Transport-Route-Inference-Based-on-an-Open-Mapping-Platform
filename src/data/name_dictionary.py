from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from src.domain.node_registry import ExternalAliasRule


NAME_DICTIONARY_FILE = "名称字典.xlsx"
NAME_DICTIONARY_HEADERS = ("节点类型", "节点全称", "节点简称")


class NameDictionaryError(ValueError):
    """Raised when the local node-name dictionary cannot be read safely."""


@dataclass(frozen=True)
class NameDictionaryEntry:
    row_number: int
    node_type: str
    full_name: str | None
    short_name: str | None

    @property
    def has_explicit_alias_mapping(self) -> bool:
        return bool(self.full_name and self.short_name)


def find_name_dictionary_file(data_dir: Path) -> Path | None:
    matches = sorted(data_dir.rglob(NAME_DICTIONARY_FILE))
    if len(matches) > 1:
        joined = "; ".join(str(path) for path in matches)
        raise NameDictionaryError(f"DATA_DIR 下存在多个 {NAME_DICTIONARY_FILE}：{joined}")
    return matches[0] if matches else None


def load_name_dictionary_entries(path: Path) -> tuple[NameDictionaryEntry, ...]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment-specific dependency error
        raise NameDictionaryError("读取名称字典需要本地 openpyxl 依赖。") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        header = tuple(_clean_text(value) for value in next(worksheet.iter_rows(values_only=True)))
        if header[:3] != NAME_DICTIONARY_HEADERS:
            expected = "、".join(NAME_DICTIONARY_HEADERS)
            actual = "、".join(value or "<空>" for value in header[:3])
            raise NameDictionaryError(
                f"{path.name} 前三列必须为 {expected}；实际为 {actual}。"
            )

        entries: list[NameDictionaryEntry] = []
        for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            first_three = tuple(row[:3]) + (None,) * max(0, 3 - len(row))
            node_type, full_name, short_name = (_clean_text(value) for value in first_three[:3])
            if not any((node_type, full_name, short_name)):
                continue
            entries.append(
                NameDictionaryEntry(
                    row_number=row_number,
                    node_type=node_type or "未分类",
                    full_name=full_name or None,
                    short_name=short_name or None,
                )
            )
        return tuple(entries)
    finally:
        workbook.close()


def build_external_alias_rules(entries: Iterable[NameDictionaryEntry]) -> tuple[ExternalAliasRule, ...]:
    """Return only explicit dictionary pairs; blank full-name rows remain review evidence."""
    pairs: list[tuple[str, str, int]] = []
    for entry in entries:
        if not entry.has_explicit_alias_mapping:
            continue
        assert entry.full_name is not None
        assert entry.short_name is not None
        if entry.full_name == entry.short_name:
            continue
        pairs.append((entry.full_name, entry.short_name, entry.row_number))

    short_to_full_names: dict[str, set[str]] = {}
    for full_name, short_name, _ in pairs:
        short_to_full_names.setdefault(short_name, set()).add(full_name)

    rules_by_full_name: dict[str, list[tuple[str, int]]] = {}
    for full_name, short_name, row_number in pairs:
        if len(short_to_full_names[short_name]) != 1:
            continue
        if full_name in short_to_full_names and short_to_full_names[full_name] != {full_name}:
            continue
        rules_by_full_name.setdefault(full_name, []).append((short_name, row_number))

    rules: list[ExternalAliasRule] = []
    for full_name, aliases_with_rows in sorted(rules_by_full_name.items()):
        aliases = tuple(sorted({alias for alias, _ in aliases_with_rows}))
        source_rows = ",".join(str(row) for _, row in aliases_with_rows)
        rules.append(
            ExternalAliasRule(
                canonical_name=full_name,
                aliases=aliases,
                source=f"{NAME_DICTIONARY_FILE}#{source_rows}",
            )
        )
    return tuple(rules)


def _clean_text(value: object) -> str:
    return "" if value is None else str(value).strip()
