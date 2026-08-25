from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path


PORT_NODE_MAINTENANCE_FILE = "码头信息维护0729版.xlsx"
PORT_NODE_MAINTENANCE_HEADERS = (
    "节点类型",
    "节点性质",
    "经营部",
    "全称",
    "简称1",
    "简称2",
    "简称3",
    "省/自治区/直辖市",
    "地级市",
    "区/县/镇",
    "街道",
    "详细地址",
    "经度",
    "纬度",
)


class PortNodeMaintenanceError(ValueError):
    """Raised when the maintained port-node workbook cannot be used safely."""


@dataclass(frozen=True)
class PortNodeMaintenanceEntry:
    row_number: int
    full_name: str
    aliases: tuple[str, ...]
    longitude: float
    latitude: float
    province: str | None
    city: str | None
    district: str | None
    address: str | None
    source: str

    @property
    def all_names(self) -> tuple[str, ...]:
        return (self.full_name, *self.aliases)


def find_port_node_maintenance_file(data_dir: Path) -> Path | None:
    matches = sorted(data_dir.rglob(PORT_NODE_MAINTENANCE_FILE))
    if len(matches) > 1:
        joined = "; ".join(str(path) for path in matches)
        raise PortNodeMaintenanceError(
            f"DATA_DIR 下存在多个 {PORT_NODE_MAINTENANCE_FILE}：{joined}"
        )
    return matches[0] if matches else None


def load_port_node_maintenance_entries(
    path: Path,
) -> tuple[PortNodeMaintenanceEntry, ...]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment-specific dependency error
        raise PortNodeMaintenanceError("读取码头信息维护表需要本地 openpyxl 依赖。") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        header = tuple(
            _clean_text(value) for value in next(worksheet.iter_rows(values_only=True))
        )
        if header[: len(PORT_NODE_MAINTENANCE_HEADERS)] != PORT_NODE_MAINTENANCE_HEADERS:
            expected = "、".join(PORT_NODE_MAINTENANCE_HEADERS)
            actual = "、".join(
                value or "<空>"
                for value in header[: len(PORT_NODE_MAINTENANCE_HEADERS)]
            )
            raise PortNodeMaintenanceError(
                f"{path.name} 表头必须为 {expected}；实际为 {actual}。"
            )

        entries: list[PortNodeMaintenanceEntry] = []
        seen_full_names: set[str] = set()
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            values = tuple(row[: len(PORT_NODE_MAINTENANCE_HEADERS)]) + (None,) * max(
                0,
                len(PORT_NODE_MAINTENANCE_HEADERS) - len(row),
            )
            (
                node_type,
                node_nature,
                _business_unit,
                full_name,
                short_name_1,
                short_name_2,
                short_name_3,
                province,
                city,
                district,
                street,
                detailed_address,
                longitude,
                latitude,
            ) = values[: len(PORT_NODE_MAINTENANCE_HEADERS)]
            if not any(values):
                continue

            clean_full_name = _required_text(full_name, path, row_number, "全称")
            if _required_text(node_type, path, row_number, "节点类型") != "站点":
                raise PortNodeMaintenanceError(
                    f"{path.name} 第 {row_number} 行节点类型不是“站点”。"
                )
            if _required_text(node_nature, path, row_number, "节点性质") != "港口码头":
                raise PortNodeMaintenanceError(
                    f"{path.name} 第 {row_number} 行节点性质不是“港口码头”。"
                )
            if clean_full_name in seen_full_names:
                raise PortNodeMaintenanceError(
                    f"{path.name} 存在重复全称：{clean_full_name}"
                )
            seen_full_names.add(clean_full_name)

            aliases = tuple(
                dict.fromkeys(
                    alias
                    for alias in (
                        _clean_text(short_name_1),
                        _clean_text(short_name_2),
                        _clean_text(short_name_3),
                    )
                    if alias and alias != clean_full_name
                )
            )
            street_text = _clean_text(street)
            detailed_text = _clean_text(detailed_address)
            address = "，".join(value for value in (street_text, detailed_text) if value)
            entries.append(
                PortNodeMaintenanceEntry(
                    row_number=row_number,
                    full_name=clean_full_name,
                    aliases=aliases,
                    longitude=_coordinate(longitude, path, row_number, "经度", -180, 180),
                    latitude=_coordinate(latitude, path, row_number, "纬度", -90, 90),
                    province=_optional_text(province),
                    city=_optional_text(city),
                    district=_optional_text(district),
                    address=address or None,
                    source=f"{path.name}#{row_number}",
                )
            )
        return tuple(entries)
    finally:
        workbook.close()


def _required_text(
    value: object,
    path: Path,
    row_number: int,
    field_name: str,
) -> str:
    text = _clean_text(value)
    if not text:
        raise PortNodeMaintenanceError(
            f"{path.name} 第 {row_number} 行缺少必填字段：{field_name}"
        )
    return text


def _optional_text(value: object) -> str | None:
    return _clean_text(value) or None


def _clean_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _coordinate(
    value: object,
    path: Path,
    row_number: int,
    field_name: str,
    minimum: float,
    maximum: float,
) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise PortNodeMaintenanceError(
            f"{path.name} 第 {row_number} 行 {field_name} 不是有效数值：{value}"
        ) from exc
    if not math.isfinite(number) or not minimum <= number <= maximum:
        raise PortNodeMaintenanceError(
            f"{path.name} 第 {row_number} 行 {field_name} 超出有效范围：{value}"
        )
    return number
